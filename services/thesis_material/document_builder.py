"""三类论文材料的通用 DOCX 构建器。"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from schemas.thesis_material import ReferenceRecord

# Match the supplied school templates: Chinese body text uses Songti, headings
# use Heiti, and Latin characters/numbers use Times New Roman.
FONT_CN = "宋体"
FONT_HEADING = "黑体"
FONT_LATIN = "Times New Roman"


def build_thesis_material_document(
    *,
    document_type: str,
    title: str,
    request: dict[str, Any],
    result: dict[str, Any],
    references: list[ReferenceRecord],
    output_path: Path,
) -> Path:
    """按文档类型构建通用 Word 文件。"""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    if document_type == "proposal_report":
        doc = _build_proposal(title, request, result, references)
    elif document_type == "literature_review":
        doc = _build_literature_review(title, request, result, references)
    elif document_type == "task_book":
        doc = _build_task_book(title, request, result, references)
    else:
        raise ValueError(f"不支持的文档类型: {document_type}")
    doc.save(str(output_path))
    return output_path


def _new_document() -> DocumentObject:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(12)
    _set_font_mapping(normal._element.get_or_add_rPr(), FONT_CN)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(0)
    for style_name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        style = doc.styles[style_name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        _set_font_mapping(style._element.get_or_add_rPr(), FONT_HEADING)
    return doc


def _build_proposal(
    title: str,
    request: dict[str, Any],
    result: dict[str, Any],
    references: list[ReferenceRecord],
) -> DocumentObject:
    doc = _new_document()
    profile = _profile_with_placeholders(request)
    context = request.get("research_context") or {}
    _cover_title(doc, profile["school"], "毕业设计（论文）开题报告", title)
    cover_fields = [
        ("课题类别", context.get("topic_category") or "某某类"),
        ("学生姓名", profile["name"]),
        ("学号", profile["student_no"]),
        ("班级", profile["class_name"]),
        ("专业", profile["major"]),
        ("指导教师", profile["internal_advisor"]),
        ("年月", profile["year_month"]),
    ]
    for label, value in cover_fields:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(4)
        paragraph.paragraph_format.space_after = Pt(10)
        paragraph.add_run(f"{label}：").bold = True
        paragraph.add_run(str(value))
    doc.add_page_break()

    table = doc.add_table(rows=0, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for heading, body in (
        ("一、本课题设计（研究）的目的", result["research_purpose"]),
        ("二、设计（研究）现状和发展趋势（文献综述）", result["research_status_and_trends"]),
        ("三、设计（研究）的主要内容", str(result.get("research_content") or "")),
        ("四、设计（研究）的重点与难点，拟采用的途径（研究手段）", _proposal_methods(result)),
        ("五、可行性分析与创新点", str(result.get("feasibility_and_innovation") or "")),
        ("六、论文（设计）写作提纲", _outline_text(result.get("writing_outline", []))),
        ("七、设计（研究）进度计划", _schedule_text(result.get("schedule", []))),
        ("八、参考文献", "\n".join(item.formatted for item in references)),
        ("指导教师意见", "\n\n\n签名：________________    年____月____日"),
        ("教研室（学术小组）意见", "\n\n\n负责人（签章）：________________    年____月____日"),
    ):
        cell = table.add_row().cells[0]
        _set_cell_width(cell, Cm(16))
        _set_cell_content(cell, heading, body)
    return doc


def _build_literature_review(
    title: str,
    request: dict[str, Any],
    result: dict[str, Any],
    references: list[ReferenceRecord],
) -> DocumentObject:
    doc = _new_document()
    _document_title(doc, title)
    _labeled_paragraph(doc, "摘要", result["abstract"])
    _labeled_paragraph(doc, "关键词", "；".join(result["keywords"]))
    sections = [
        ("一、引言", result["introduction"]),
        ("二、国内研究现状", result["domestic_research"]),
        ("三、国外研究现状", result["foreign_research"]),
    ]
    for heading, body in sections:
        doc.add_heading(heading, level=1)
        _body_paragraph(doc, body)
    doc.add_heading("四、主题分类与代表性研究", level=1)
    for index, theme in enumerate(result.get("themes", []), start=1):
        doc.add_heading(f"4.{index} {theme.get('title', '')}", level=2)
        _body_paragraph(doc, str(theme.get("content") or ""))
    for heading, key in (
        ("五、研究方法比较", "method_comparison"),
        ("六、现有研究不足", "research_gaps"),
        ("七、发展趋势", "future_trends"),
        ("八、结论", "conclusion"),
    ):
        doc.add_heading(heading, level=1)
        _body_paragraph(doc, result[key])
    doc.add_heading("参考文献", level=1)
    for item in references:
        paragraph = doc.add_paragraph(item.formatted)
        paragraph.paragraph_format.first_line_indent = Cm(-0.74)
        paragraph.paragraph_format.left_indent = Cm(0.74)
        paragraph.paragraph_format.line_spacing = 1.25
    return doc


def _build_task_book(
    title: str,
    request: dict[str, Any],
    result: dict[str, Any],
    references: list[ReferenceRecord],
) -> DocumentObject:
    doc = _new_document()
    profile = _profile_with_placeholders(request)
    _cover_title(doc, profile["school"], "毕业设计任务书", "")
    info = doc.add_table(rows=6, cols=6)
    info.style = "Table Grid"
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_rows = [
        ("二级学院", profile["college"], "姓名", profile["name"], "校内指导教师", profile["internal_advisor"]),
        ("班级名称", profile["class_name"], "学号", profile["student_no"], "企业指导教师", profile["enterprise_advisor"]),
    ]
    for row_index, row_values in enumerate(info_rows):
        for column, value in enumerate(row_values):
            _replace_cell_text(info.cell(row_index, column), str(value), bold=column % 2 == 0)
    _merge_labeled_row(info, 2, "选题名称", title)
    _merge_labeled_row(info, 3, "选题类型", str(request.get("topic_type") or "产品设计类"))
    goals = "\n".join(f"{index}. {item}" for index, item in enumerate(result.get("design_goals", []), start=1))
    stack = "、".join(str(item) for item in result.get("technology_stack", []))
    design_goal = f"{result.get('design_background', '')}\n技术建议：{stack}\n{goals}".strip()
    _merge_labeled_row(info, 4, "设计目标", design_goal)
    task_lines = []
    for item in result.get("module_tasks", []):
        task_lines.append(
            f"{item.get('name', '模块')}：面向{item.get('role', '相关用户')}，{item.get('responsibilities', '')}；边界：{item.get('boundary', '')}"
        )
    _merge_labeled_row(info, 5, "设计任务", "\n".join(task_lines))

    schedule_items = result.get("schedule_items", [])
    schedule = doc.add_table(rows=1, cols=5)
    schedule.style = "Table Grid"
    schedule.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, heading in enumerate(("序号", "设计任务", "起始时间", "结束时间", "阶段成果")):
        _replace_cell_text(schedule.cell(0, index), heading, bold=True)
    for index, item in enumerate(schedule_items, start=1):
        cells = schedule.add_row().cells
        schedule_values = (
            str(index),
            str(item.get("task") or ""),
            str(item.get("start") or ""),
            str(item.get("end") or ""),
            str(item.get("deliverable") or ""),
        )
        for column, value in enumerate(schedule_values):
            _replace_cell_text(cells[column], value)

    indicators = "\n".join(
        f"（{index}）{value}" for index, value in enumerate(result.get("main_indicators", []), start=1)
    )
    references_text = "\n".join(item.formatted for item in references)
    outcome = doc.add_table(rows=6, cols=2)
    outcome.style = "Table Grid"
    outcome.alignment = WD_TABLE_ALIGNMENT.CENTER
    _replace_cell_text(outcome.cell(0, 0), "预期成果", bold=True)
    forms = "\n".join(f"（{index}）{value}" for index, value in enumerate(result.get("deliverable_forms", []), start=1))
    requirements = "\n".join(
        f"（{index}）{value}" for index, value in enumerate(result.get("deliverable_requirements", []), start=1)
    )
    _replace_cell_text(outcome.cell(0, 1), f"成果表现形式\n{forms}\n成果要求\n{requirements}")
    _replace_cell_text(outcome.cell(1, 0), "主要指标", bold=True)
    _replace_cell_text(outcome.cell(1, 1), indicators)
    _replace_cell_text(outcome.cell(2, 0), "主要参考资料", bold=True)
    _replace_cell_text(outcome.cell(2, 1), references_text)
    for row, label in ((3, "指导教师"), (4, "教研室审核意见"), (5, "二级学院审核意见")):
        _replace_cell_text(outcome.cell(row, 0), label, bold=True)
        _replace_cell_text(outcome.cell(row, 1), "\n（签名）________________    年____月____日")
    note = doc.add_paragraph("注：⑴ 请双面打印。⑵ 如需附图，请以附件形式提供。")
    note.paragraph_format.space_before = Pt(8)
    return doc


def _cover_title(doc: DocumentObject, school: str, document_name: str, title: str) -> None:
    if school:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(school)
        _format_run(run, 18, bold=True, font=FONT_HEADING)
        paragraph.paragraph_format.space_after = Pt(30)
    _document_title(doc, document_name)
    if title:
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(45)
        paragraph.paragraph_format.space_after = Pt(40)
        run = paragraph.add_run(title)
        _format_run(run, 16, bold=True, font=FONT_HEADING)


def _document_title(doc: DocumentObject, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(20)
    run = paragraph.add_run(text)
    _format_run(run, 22, bold=True, font=FONT_HEADING)


def _labeled_paragraph(doc: DocumentObject, label: str, body: str) -> None:
    paragraph = doc.add_paragraph()
    label_run = paragraph.add_run(f"{label}：")
    label_run.bold = True
    paragraph.add_run(_normalize_generated_text(body))
    paragraph.paragraph_format.first_line_indent = Cm(0.74)


def _body_paragraph(doc: DocumentObject, body: str) -> None:
    for block in [item.strip() for item in _normalize_generated_text(body).split("\n") if item.strip()]:
        paragraph = doc.add_paragraph(block)
        paragraph.paragraph_format.first_line_indent = Cm(0.74)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY


def _proposal_methods(result: dict[str, Any]) -> str:
    return "\n".join(
        (
            "（一）研究重点\n" + result["key_points"],
            "（二）研究难点\n" + result["difficulties"],
            "（三）拟采用的途径与研究手段\n" + result["research_methods"],
        )
    )


def _schedule_text(items: list[dict[str, Any]]) -> str:
    return "\n".join(
        f"{item.get('start', '')}-{item.get('end', '')}：{item.get('task', '')}；阶段成果：{item.get('deliverable', '')}"
        for item in items
    )


def _outline_text(items: list[dict[str, Any]]) -> str:
    """把模型生成的结构化论文提纲转换成三级编号文本。"""

    lines: list[str] = []
    for chapter_index, item in enumerate(items, start=1):
        lines.append(f"{chapter_index} {_strip_outline_number(str(item.get('title') or ''))}")
        seen_section_titles: set[str] = set()
        raw_sections = item.get("sections")
        sections: list[Any] = raw_sections if isinstance(raw_sections, list) else []
        for section_index, section in enumerate(sections, start=1):
            if isinstance(section, dict):
                section_title = _strip_outline_number(str(section.get("title") or ""))
                raw_subsection_items = section.get("subsections")
                if not isinstance(raw_subsection_items, list):
                    raw_subsection_items = section.get("items")
                subsection_items: list[Any] = raw_subsection_items if isinstance(raw_subsection_items, list) else []
            else:
                section_title = _strip_outline_number(str(section))
                subsection_items = []
            if not section_title or section_title in seen_section_titles:
                continue
            seen_section_titles.add(section_title)
            lines.append(f"{chapter_index}.{section_index} {section_title}")
            for subsection_index, subsection in enumerate(subsection_items, start=1):
                lines.append(
                    f"{chapter_index}.{section_index}.{subsection_index} {_strip_outline_number(str(subsection))}"
                )
        raw_subsections = item.get("subsections")
        subsections: list[Any] = raw_subsections if isinstance(raw_subsections, list) else []
        for section_index, section in enumerate(subsections, start=len(sections) + 1):
            if not isinstance(section, dict):
                continue
            section_title = _strip_outline_number(str(section.get("title") or ""))
            if not section_title or section_title in seen_section_titles:
                continue
            seen_section_titles.add(section_title)
            lines.append(
                f"{chapter_index}.{section_index} {section_title}"
            )
            raw_third_level_items = section.get("items")
            third_level_items: list[Any] = raw_third_level_items if isinstance(raw_third_level_items, list) else []
            for subsection_index, subsection in enumerate(third_level_items, start=1):
                lines.append(
                    f"{chapter_index}.{section_index}.{subsection_index} {_strip_outline_number(str(subsection))}"
                )
    return "\n".join(lines)


def _strip_outline_number(value: str) -> str:
    """移除模型偶尔自带的章节编号，避免与程序编号叠加。"""

    stripped = re.sub(r"^\s*(?:\d+(?:\.\d+)*[、.．\s]+)+", "", value)
    stripped = re.sub(r"^\s*第\s*\d+\s*[章节篇]\s*", "", stripped)
    stripped = re.sub(r"^\s*第?[一二三四五六七八九十]+[章节、.．\s]+", "", stripped)
    return stripped.strip()


def _normalize_generated_text(value: str) -> str:
    """清理模型正文中常见的中文标点空格和重复分隔符。"""

    normalized = re.sub(r"\s+([，。；：！？、）】])", r"\1", value)
    normalized = re.sub(r"([（【])\s+", r"\1", normalized)
    normalized = normalized.replace("。；", "；").replace("；。", "；")
    return normalized


def _profile_with_placeholders(request: dict[str, Any]) -> dict[str, str]:
    """为非必填身份字段补充可见且易替换的通用占位值。"""

    raw_profile = request.get("student_profile")
    profile: dict[str, Any] = raw_profile if isinstance(raw_profile, dict) else {}
    defaults = {
        "school": "某某大学",
        "college": "某某学院",
        "name": "某某某",
        "student_no": "20XXXXXXXXXX",
        "class_name": "某某班",
        "major": "某某专业",
        "internal_advisor": "某某某",
        "enterprise_advisor": "某某某",
        "year_month": "20XX年XX月",
    }
    return {key: str(profile.get(key) or value) for key, value in defaults.items()}


def _set_cell_content(cell: Any, heading: str, body: str) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    heading_paragraph = cell.paragraphs[0]
    heading_run = heading_paragraph.add_run(heading)
    _format_run(heading_run, 12, bold=True, font=FONT_HEADING)
    for block in [item.strip() for item in _normalize_generated_text(body).split("\n")]:
        paragraph = cell.add_paragraph(block)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.first_line_indent = Cm(0.74) if block else None
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for run in paragraph.runs:
            _format_run(run, 12)
    _set_cell_margin(cell, 120)


def _merge_labeled_row(table: Any, row: int, label: str, content: str) -> None:
    _replace_cell_text(table.cell(row, 0), label, bold=True)
    merged = table.cell(row, 1).merge(table.cell(row, 5))
    _replace_cell_text(merged, content)


def _replace_cell_text(cell: Any, text: str, *, bold: bool = False) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    lines = _normalize_generated_text(text).split("\n") or [""]
    for index, line in enumerate(lines):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(line) < 20 else WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.line_spacing = 1.25
        run = paragraph.add_run(line)
        _format_run(run, 10.5, bold=bold)
    _set_cell_margin(cell, 100)


def _format_run(run: Any, size: float, *, bold: bool = False, font: str = FONT_CN) -> None:
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.bold = bold
    _set_font_mapping(run._element.get_or_add_rPr(), font)


def _set_font_mapping(rpr: Any, east_asia: str) -> None:
    """Write explicit Word font slots so CJK and Latin text are stable."""

    r_fonts = rpr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        rpr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:cs"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), east_asia)


def _set_cell_margin(cell: Any, value: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_width(cell: Any, width: Cm) -> None:
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_width = tc_pr.first_child_found_in("w:tcW")
    if tc_width is None:
        tc_width = OxmlElement("w:tcW")
        tc_pr.append(tc_width)
    tc_width.set(qn("w:w"), str(int(width.twips)))
    tc_width.set(qn("w:type"), "dxa")


__all__ = ["build_thesis_material_document"]
