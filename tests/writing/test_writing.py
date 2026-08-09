import asyncio
from pathlib import Path

from docx import Document
from pydantic import ValidationError

from app import app
from schemas.writing import (
    LiteratureReviewRequest,
    ProposalReportRequest,
    ReferenceRecord,
    TaskBookRequest,
)
from services.writing.document_builder import build_writing_document
from services.writing.generation import (
    PROPOSAL_SCHEDULE,
    TASK_BOOK_SCHEDULE,
    _build_schedule,
    _validate_result,
)
from services.writing.llm_service import (
    LITERATURE_BODY_LENGTH,
    _literature_body_length,
    _trim_to_complete_sentences,
    repair_length_constraints,
)
from services.writing.reference_service import (
    _merge_records,
    _select_records,
    _target_language_quota,
    parse_reference_records,
)


def test_writing_routes_are_registered() -> None:
    routes = {route.path for route in app.routes}
    assert "/api/v1/writing/products" in routes
    assert "/api/v1/writing/proposal-reports" in routes
    assert "/api/v1/writing/literature-reviews" in routes
    assert "/api/v1/writing/task-books" in routes
    assert "/api/v1/writing/tasks/{task_id}" in routes
    assert "/api/v1/writing/tasks/{task_id}/events" in routes
    assert "/api/v1/writing/tasks/{task_id}/download" in routes


def test_title_is_only_required_request_field() -> None:
    proposal = ProposalReportRequest(title="智慧校园管理平台的设计与实现")
    review = LiteratureReviewRequest(title="生成式人工智能教育应用研究综述")
    task_book = TaskBookRequest(title="校园饭卡管理系统的设计与实现")
    assert proposal.reference_options.target_count == 15
    assert review.reference_options.target_count == 20
    assert task_book.student_profile.name is None


def test_schedule_rejects_reversed_dates() -> None:
    try:
        ProposalReportRequest(
            title="智慧校园管理平台的设计与实现",
            schedule_options={"start_date": "2026-05-01", "end_date": "2026-04-01"},
        )
    except ValidationError as exc:
        assert "开始日期不能晚于结束日期" in str(exc)
    else:
        raise AssertionError("倒序日期应校验失败")


def test_parse_reference_records_renumbers_and_deduplicates() -> None:
    text = "\n".join(
        (
            "[1]张三,李四.智慧校园平台研究[J].软件导刊,2024,12(3):10-18.",
            "[2]Smith J.Cloud Campus Architecture[J].Computing Review,2023,8(2):20-30.",
            "[3]张三,李四.智慧校园平台研究[J].软件导刊,2024,12(3):10-18.",
        )
    )
    records = parse_reference_records(text, provider="mixed")
    assert len(records) == 2
    assert records[0].index == 1
    assert records[0].title == "智慧校园平台研究"
    assert records[1].language == "en"


def test_merge_reference_records_deduplicates_and_renumbers() -> None:
    first = [ReferenceRecord(index=1, title="文献甲", formatted="[1]作者.文献甲[J].期刊,2024.")]
    second = [
        ReferenceRecord(index=1, title="文献甲", formatted="[1]作者.文献甲[J].期刊,2024."),
        ReferenceRecord(index=2, title="文献乙", formatted="[2]作者.文献乙[J].期刊,2023."),
    ]
    merged = _merge_records(first, second, 10)
    assert [item.index for item in merged] == [1, 2]
    assert merged[1].formatted.startswith("[2]")


def test_reference_language_quota_and_selection() -> None:
    assert _target_language_quota(15) == (10, 5)
    assert _target_language_quota(20) == (14, 6)
    records = [
        ReferenceRecord(index=index, title=f"中文{index}", language="zh", formatted=f"[{index}]中文{index}")
        for index in range(1, 11)
    ] + [
        ReferenceRecord(index=index + 10, title=f"English {index}", language="en", formatted=f"[{index + 10}]English {index}")
        for index in range(1, 6)
    ]
    selected = _select_records(records, 15, 10, 5)
    assert len(selected) == 15
    assert sum(item.language == "zh" for item in selected) == 10
    assert sum(item.language == "en" for item in selected) == 5
    assert [item.index for item in selected] == list(range(1, 16))


def test_reference_validation_requires_every_reference_to_be_cited() -> None:
    result = {
        "abstract": "摘" * 220,
        "introduction": "已有研究[1]。",
        "references": [{"index": 1}, {"index": 2}],
    }
    try:
        _validate_result("literature_review", result, 2)
    except RuntimeError as exc:
        assert "未在正文引用" in str(exc)
    else:
        raise AssertionError("未引用全部文献时应校验失败")


def test_length_repair_only_regenerates_invalid_fields(monkeypatch) -> None:
    calls: list[str] = []

    async def fake_ask_text(system: str, prompt: str, *, max_tokens: int = 5000) -> str:
        del system, max_tokens
        calls.append(prompt)
        return "修" * 850

    monkeypatch.setattr("services.writing.llm_service._ask_text", fake_ask_text)
    result = {
        "research_purpose": "长" * 1200,
        "research_status_and_trends": "现" * 1600,
        "key_points": "重" * 350,
        "difficulties": "难" * 350,
        "research_methods": "法" * 450,
    }

    asyncio.run(repair_length_constraints("proposal_report", {"title": "测试课题"}, result))

    assert len(calls) == 1
    assert len(result["research_purpose"]) == 850
    assert len(result["research_status_and_trends"]) == 1600


def test_literature_total_length_is_trimmed_without_regenerating_sections(monkeypatch) -> None:
    async def unexpected_llm_call(*_args: object, **_kwargs: object) -> str:
        raise AssertionError("完整句裁剪足够时不应调用模型")

    def fixed_length_sentence(prefix: str, fill: str, length: int) -> str:
        return prefix + fill * (length - len(prefix) - 1) + "。"

    monkeypatch.setattr("services.writing.llm_service._ask_text", unexpected_llm_call)
    theme = "".join(
        [fixed_length_sentence("已有做法、观点、优势、不足和小结[1]", "甲", 100)]
        + [fixed_length_sentence("比较分析", "乙", 100) for _ in range(6)]
    )
    result = {
        "abstract": fixed_length_sentence("摘要", "甲", 300),
        "introduction": fixed_length_sentence("引言", "甲", 600),
        "domestic_research": fixed_length_sentence("国内研究", "甲", 800),
        "foreign_research": fixed_length_sentence("国外研究", "甲", 800),
        "themes": [{"title": f"主题{index}", "content": theme} for index in range(5)],
        "method_comparison": fixed_length_sentence("方法比较", "甲", 700),
        "research_gaps": fixed_length_sentence("研究不足", "甲", 500),
        "future_trends": fixed_length_sentence("未来趋势", "甲", 500),
        "conclusion": fixed_length_sentence("结论", "甲", 450),
    }

    asyncio.run(repair_length_constraints("literature_review", {"title": "测试课题"}, result))

    assert _literature_body_length(result) <= LITERATURE_BODY_LENGTH[1]
    assert all(500 <= len(item["content"]) <= 700 for item in result["themes"])


def test_sentence_trim_preserves_citations_and_required_topics() -> None:
    text = "".join(
        (
            "行业背景" + "甲" * 120 + "。",
            "现实问题" + "乙" * 120 + "。",
            "技术背景" + "丙" * 120 + "。",
            "研究必要性" + "丁" * 120 + "。",
            "应用价值" + "戊" * 120 + "。",
            "研究目标" + "己" * 120 + "。",
            "补充分析" + "庚" * 120 + "。",
            "相关依据[1]" + "辛" * 120 + "。",
            "重复说明" + "壬" * 120 + "。",
        )
    )

    trimmed = _trim_to_complete_sentences(text, "research_purpose", 700, 1000)

    assert 700 <= len(trimmed) <= 1000
    assert "[1]" in trimmed
    for keyword in ("背景", "问题", "技术", "必要性", "价值", "目标"):
        assert keyword in trimmed


def test_relative_schedule_has_required_number_of_stages() -> None:
    proposal = _build_schedule({}, PROPOSAL_SCHEDULE, default_weeks=16)
    task_book = _build_schedule({}, TASK_BOOK_SCHEDULE, default_weeks=20)
    assert len(proposal) == 9
    assert proposal[0]["start"] == "第1周"
    assert proposal[-1]["end"] == "第16周"
    assert len(task_book) == 5


def test_build_three_document_types(tmp_path: Path) -> None:
    references = [
        ReferenceRecord(index=1, title="测试文献", authors=["张三"], year="2024", formatted="[1]张三.测试文献[J].测试期刊,2024,1(1):1-5.")
    ]
    request = {"student_profile": {}, "research_context": {}}
    proposal_result = {
        "research_purpose": "研究目的[1]。",
        "research_status_and_trends": "研究现状[1]。",
        "key_points": "研究重点。",
        "difficulties": "研究难点。",
        "research_methods": "研究方法。",
        "schedule": _build_schedule({}, PROPOSAL_SCHEDULE, default_weeks=16),
    }
    review_result = {
        "abstract": "摘要。",
        "keywords": ["测试", "综述", "方法"],
        "introduction": "引言[1]。",
        "domestic_research": "国内研究[1]。",
        "foreign_research": "国外研究[1]。",
        "themes": [{"title": "主题", "content": "比较研究[1]。"}] * 3,
        "method_comparison": "方法比较。",
        "research_gaps": "不足。",
        "future_trends": "趋势。",
        "conclusion": "结论。",
    }
    task_result = {
        "design_background": "设计背景。",
        "technology_stack": ["Java", "Vue"],
        "design_goals": [f"目标{i}" for i in range(1, 6)],
        "module_tasks": [
            {"name": f"模块{i}", "role": "用户", "responsibilities": "完成业务", "boundary": "仅处理本模块"}
            for i in range(1, 5)
        ],
        "schedule_items": _build_schedule({}, TASK_BOOK_SCHEDULE, default_weeks=20),
        "deliverable_forms": ["成果文档", "项目源文件"],
        "deliverable_requirements": ["项目可运行", "测试通过"],
    }
    paths = []
    for document_type, result in (
        ("proposal_report", proposal_result),
        ("literature_review", review_result),
        ("task_book", task_result),
    ):
        path = tmp_path / f"{document_type}.docx"
        build_writing_document(
            document_type=document_type,
            title="测试课题",
            request=request,
            result=result,
            references=references if document_type != "task_book" else [],
            output_path=path,
        )
        paths.append(path)
    assert all(path.exists() for path in paths)
    assert "指导教师意见" in "\n".join(cell.text for table in Document(paths[0]).tables for row in table.rows for cell in row.cells)
    assert "审核意见" in "\n".join(cell.text for table in Document(paths[2]).tables for row in table.rows for cell in row.cells)
