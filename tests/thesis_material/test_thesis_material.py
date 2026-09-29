import asyncio
import zipfile
from pathlib import Path
from types import SimpleNamespace
from typing import Any, Literal
from unittest.mock import AsyncMock

import pytest
from docx import Document
from pydantic import ValidationError

import services.thesis_material.llm_service as llm_service
from app import app
from schemas.thesis_material import (
    LiteratureReviewRequest,
    ProposalReportRequest,
    ReferenceRecord,
    TaskBookRequest,
    ThesisMaterialSubmitResponse,
)
from services.thesis_material import generation as material_generation
from services.thesis_material.document_builder import (
    _normalize_generated_text,
    _outline_text,
    build_thesis_material_document,
)
from services.thesis_material.generation import (
    PROPOSAL_SCHEDULE,
    TASK_BOOK_SCHEDULE,
    _build_schedule,
    _ensure_body_minimum_after_normalization,
    _ensure_reference_coverage_after_normalization,
    _material_body_char_count,
    _missing_reference_indexes,
    _validate_result,
    _word_count_metadata,
    prepare_material_plan,
)
from services.thesis_material.llm_service import (
    _literature_body_length,
    _mark_unconfirmed_generated_metrics,
    _normalize_keywords,
    _repair_missing_text_fields,
    _trim_to_character_limit,
    _trim_to_complete_sentences,
    build_length_plan,
    normalize_citation_claims,
    repair_length_constraints,
    repair_reference_coverage,
)
from services.thesis_material.order_service import ThesisMaterialOrderService
from services.thesis_material.profile_policy import (
    missing_profile_fields,
    profile_with_placeholders,
)
from services.thesis_material.reference_service import (
    _merge_records,
    _rank_records_by_relevance,
    _select_records,
    parse_reference_records,
)


def test_literature_keywords_accept_delimited_text_and_outline_fallback() -> None:
    outline = [
        {"chapter": "研究背景与意义", "sections": [{"name": "校园二手交易业务流程"}]},
        {"chapter": "平台架构设计", "sections": [{"name": "前后端分离架构"}]},
    ]

    assert _normalize_keywords(
        "校园二手交易；Spring Boot、Vue",
        title="校园二手交易平台研究",
        source_outline=outline,
    ) == ["校园二手交易", "Spring Boot", "Vue"]
    assert _normalize_keywords(
        None,
        title="校园二手交易平台研究",
        source_outline=outline,
    ) == ["校园二手交易平台研究", "校园二手交易业务流程", "前后端分离架构"]


def test_literature_missing_text_field_is_repaired(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[str] = []

    async def fake_ask_json(system: str, prompt: str, *, max_tokens: int = 5000) -> dict[str, str]:
        calls.append(prompt)
        return {"research_gaps": "现有研究仍缺少跨场景的数据验证。"}

    monkeypatch.setattr(llm_service, "_ask_json", fake_ask_json)
    plan = build_length_plan("literature_review", {"target_word_count": 3500})
    repaired = asyncio.run(
        _repair_missing_text_fields(
            {"method_comparison": "已有方法比较。", "research_gaps": ""},
            ("method_comparison", "research_gaps"),
            title="测试课题",
            reference_text="[1]测试文献。",
            length_plan=plan,
        )
    )

    assert repaired["method_comparison"] == "已有方法比较。"
    assert repaired["research_gaps"] == "现有研究仍缺少跨场景的数据验证。"
    assert len(calls) == 1
    assert "research_gaps" in str(calls[0])


def test_final_reference_coverage_uses_verifiable_title_statements() -> None:
    references = [
        ReferenceRecord(
            index=index,
            title=f"校园二手交易研究{index}",
            authors=["测试作者"],
            year="2025",
            formatted=f"[{index}]测试作者.校园二手交易研究{index}[J].测试期刊,2025(1):1-5.",
        )
        for index in range(1, 13)
    ]
    result: dict[str, Any] = {
        field: "校园交易平台相关研究。" * 40
        for field in (
            "abstract",
            "introduction",
            "domestic_research",
            "foreign_research",
            "method_comparison",
            "research_gaps",
            "future_trends",
            "conclusion",
        )
    }
    result["domestic_research"] += "".join(f"已有文献线索[{index}]。" for index in range(1, 11))
    result["themes"] = [{"title": f"主题{index}", "content": "校园二手交易研究比较。" * 40} for index in range(1, 4)]

    _ensure_reference_coverage_after_normalization(
        "literature_review",
        {"target_word_count": 3500},
        result,
        references,
    )

    assert "研究线索[11]" in result["conclusion"]
    assert "研究线索[12]" in result["conclusion"]
    assert "发现" not in result["conclusion"]


def test_thesis_material_routes_are_registered() -> None:
    routes = {getattr(route, "path", "") for route in app.routes}
    assert "/api/v1/thesis-materials/products" in routes
    assert "/api/v1/thesis-materials/proposal-reports" in routes
    assert "/api/v1/thesis-materials/literature-reviews" in routes
    assert "/api/v1/thesis-materials/task-books" in routes
    assert "/api/v1/thesis-materials/tasks/{task_id}" in routes
    assert "/api/v1/thesis-materials/tasks/{task_id}/events" in routes
    assert "/api/v1/thesis-materials/tasks/{task_id}/download" in routes
    assert "/api/v1/admin/thesis-material-orders" in routes
    assert "/api/v1/admin/thesis-material-orders/{order_id}" in routes

    # 旧版调用方升级前仍能访问隐藏兼容路由。
    assert "/api/v1/paper-materials/products" in routes
    assert "/api/v1/paper-materials/tasks/{task_id}" in routes
    assert "/api/v1/admin/paper-material-orders" in routes
    assert "/api/v1/admin/paper-material-orders/{order_id}" in routes
    assert "/api/v1/writing/products" in routes
    assert "/api/v1/writing/proposal-reports" in routes
    assert "/api/v1/writing/literature-reviews" in routes
    assert "/api/v1/writing/task-books" in routes
    assert "/api/v1/writing/tasks/{task_id}" in routes
    assert "/api/v1/writing/tasks/{task_id}/events" in routes
    assert "/api/v1/writing/tasks/{task_id}/download" in routes
    assert "/api/v1/admin/writing-orders" in routes
    assert "/api/v1/admin/writing-orders/{order_id}" in routes


def test_thesis_material_openapi_only_exposes_canonical_routes() -> None:
    paths = app.openapi()["paths"]
    assert "/api/v1/thesis-materials/products" in paths
    assert "/api/v1/admin/thesis-material-orders" in paths
    assert "/api/v1/paper-materials/products" not in paths
    assert "/api/v1/admin/paper-material-orders" not in paths
    assert "/api/v1/writing/products" not in paths
    assert "/api/v1/admin/writing-orders" not in paths


def test_outline_text_rebuilds_numbering_without_duplicates() -> None:
    outline = [
        {
            "title": "第1章 绪论",
            "sections": [
                {
                    "title": "1.1 研究背景",
                    "subsections": ["1.1.1 现实问题", "1.1.2 研究价值"],
                }
            ],
        }
    ]
    assert _outline_text(outline).splitlines() == [
        "1 绪论",
        "1.1 研究背景",
        "1.1.1 现实问题",
        "1.1.2 研究价值",
    ]


def test_outline_text_keeps_same_named_user_sections_in_original_order() -> None:
    outline = [
        {
            "title": "绪论",
            "sections": [
                {"title": "研究背景", "subsections": []},
                {"title": "研究背景", "subsections": ["不同场景"]},
            ],
        }
    ]
    assert _outline_text(outline).splitlines() == [
        "1 绪论",
        "1.1 研究背景",
        "1.2 研究背景",
        "1.2.1 不同场景",
    ]


def test_generated_text_normalizes_chinese_punctuation() -> None:
    assert _normalize_generated_text("模块 A ；边界： 不处理。 ；完成") == "模块 A；边界： 不处理；完成"


def test_material_requires_outline_and_inherits_reference_config() -> None:
    payload = {
        "title": "通用课题研究",
        "source_outline": [{"chapter": "绪论", "sections": [{"name": "研究背景", "abstract": "背景"}]}],
        "thesis_config": {"chinese_reference_count": 2, "english_reference_count": 1},
    }
    for request_type in (ProposalReportRequest, LiteratureReviewRequest, TaskBookRequest):
        request = request_type.model_validate(payload)
        assert request.thesis_config.chinese_reference_count == 2
        assert request.thesis_config.english_reference_count == 1
        assert "reference_options" not in request.model_dump()
        assert "schedule_options" not in request.model_dump()
        with pytest.raises(ValidationError):
            request_type.model_validate({"title": "通用课题"})
    assert TaskBookRequest.model_validate(payload).target_word_count == 2000


def test_material_requests_accept_confirmed_outline_and_common_config() -> None:
    outline = [
        {
            "chapter": "第1章 绪论",
            "sections": [
                {
                    "name": "1.1 研究背景",
                    "abstract": "概述研究背景",
                    "subsections": [{"name": "1.1.1 现实需求", "abstract": "说明现实需求"}],
                }
            ],
        }
    ]
    payload = {
        "title": "校园服务平台研究",
        "source_outline": outline,
        "thesis_config": {
            "target_word_count": 8000,
            "aboutmsg": "围绕校园服务流程",
        },
    }
    for request_type in (ProposalReportRequest, LiteratureReviewRequest, TaskBookRequest):
        request = request_type.model_validate(payload)
        snapshot = request.model_dump(mode="json")
        assert snapshot["source_outline"][0]["chapter"] == "绪论"
        assert "codetype" not in snapshot["thesis_config"]
    assert ProposalReportRequest.model_validate(payload).target_word_count == 4000
    assert LiteratureReviewRequest.model_validate(payload).target_word_count == 6000


def test_task_book_recognizes_user_supplied_numeric_requirements_in_common_config() -> None:
    request = {"thesis_config": {"aboutmsg": "由导师确认并发 50 人作为测试要求"}}
    assert llm_service._has_user_numeric_requirements(request) is True
    assert llm_service._has_user_numeric_requirements({"thesis_config": {"target_word_count": 8000}}) is False
    assert llm_service._has_user_numeric_requirements({"thesis_config": {"aboutmsg": "使用Spring Boot 3"}}) is False


@pytest.mark.asyncio
async def test_material_requires_confirmed_outline() -> None:
    with pytest.raises(ValueError, match="确认论文大纲"):
        await prepare_material_plan("proposal_report", {"title": "通用课题"})


def test_invalid_source_outline_is_rejected_before_paid_submission() -> None:
    with pytest.raises(ValidationError):
        ProposalReportRequest.model_validate(
            {
                "title": "校园服务平台研究",
                "source_outline": [{"chapter": "绪论", "sections": []}],
            }
        )


def test_default_material_schedule_does_not_assume_a_software_project() -> None:
    proposal = "\n".join(item[0] for item in PROPOSAL_SCHEDULE)
    task_book = "\n".join(item[0] for item in TASK_BOOK_SCHEDULE)
    for technical_phrase in ("数据库", "后端", "接口联调", "系统测试"):
        assert technical_phrase not in proposal
        assert technical_phrase not in task_book


@pytest.mark.asyncio
async def test_proposal_uses_confirmed_outline_without_regenerating_it(monkeypatch: pytest.MonkeyPatch) -> None:
    request = ProposalReportRequest.model_validate(
        {
            "title": "校园服务平台研究",
            "source_outline": [
                {
                    "chapter": "第一章 绪论",
                    "sections": [
                        {
                            "name": "1.1 研究背景",
                            "abstract": "背景",
                            "subsections": [{"name": "1.1.1 业务需求", "abstract": "需求"}],
                        }
                    ],
                }
            ],
            "thesis_config": {"aboutmsg": "聚焦校园场景"},
        }
    ).model_dump(mode="json")
    request["material_outline"] = {"sections": [{"title": "用户指定材料结构"}]}
    prompts: list[str] = []

    async def fake_ask_text(_system: str, prompt: str, *, max_tokens: int = 5000) -> str:
        del max_tokens
        prompts.append(prompt)
        return "正文"

    async def fake_ask_json(_system: str, prompt: str, *, max_tokens: int = 5000) -> dict[str, str]:
        del max_tokens
        prompts.append(prompt)
        return {
            key: "正文"
            for key in (
                "research_content",
                "key_points",
                "difficulties",
                "research_methods",
                "feasibility_and_innovation",
            )
        }

    monkeypatch.setattr(llm_service, "_ask_text", fake_ask_text)
    monkeypatch.setattr(llm_service, "_ask_json", fake_ask_json)
    result = await llm_service.generate_proposal_content(request, [])

    assert result["writing_outline"] == [
        {
            "title": "绪论",
            "sections": [{"title": "研究背景", "subsections": ["业务需求"]}],
        }
    ]
    assert "已有用户确认的论文大纲" in str(prompts[-1])
    assert "聚焦校园场景" in str(prompts[-1])
    assert "codetype" not in str(prompts[-1])
    assert "用户指定材料结构" in str(prompts[-1])


def test_submit_response_exposes_missing_profile_fields() -> None:
    fields = ThesisMaterialSubmitResponse.model_fields
    assert "missing_profile_fields" in fields


def test_missing_profile_fields_are_fixed_and_placeholders_are_not_fake_data() -> None:
    missing = missing_profile_fields("proposal_report")
    profile = profile_with_placeholders()
    assert "school" in missing
    assert "student_no" in missing
    assert profile["school"] == "【待补充：学校】"
    assert profile["student_no"] == "【待补充：学号】"


def test_thesis_material_order_number_uses_domain_prefix() -> None:
    assert ThesisMaterialOrderService.generate_order_sn().startswith("TM")


def test_schedule_rejects_reversed_dates() -> None:
    try:
        ProposalReportRequest.model_validate(
            {
                "title": "智慧校园管理平台的设计与实现",
                "schedule_options": {"start_date": "2026-05-01", "end_date": "2026-04-01"},
            }
        )
    except ValidationError as exc:
        assert "schedule_options" in str(exc)
    else:
        raise AssertionError("倒序日期应校验失败")


def test_parse_reference_records_renumbers_and_deduplicates() -> None:
    text = "\n".join(
        (
            "[1]张三,李四.智慧校园平台研究[J].软件导刊,2024,12(3):10-18.",
            "[2]Smith J.Cloud Campus Architecture[J].Computing Review,2023,8(2):20-30.",
            "[3]张三,李四.智慧校园平台研究[J].软件导刊,2024,12(3):10-18.",
        )
    )
    records = parse_reference_records(text, provider="mixed")
    assert len(records) == 2
    assert records[0].index == 1
    assert records[0].title == "智慧校园平台研究"
    assert records[0].authors == ["张三", "李四"]
    assert records[1].language == "en"


def test_parse_reference_title_preserves_dot_in_technology_name() -> None:
    records = parse_reference_records(
        "[1]曲蕴慧.基于ASP.NET的校园二手交易平台设计与实现[J].数字技术与应用,2013,(7):110-111.",
        provider="wfapi",
    )
    assert records[0].authors == ["曲蕴慧"]
    assert records[0].title == "基于ASP.NET的校园二手交易平台设计与实现"


def test_parse_reference_records_rejects_truncated_title() -> None:
    """带省略号的检索片段不是完整题名，不得进入最终文献。"""

    records = parse_reference_records(
        "[1]Smith J.The effects of second-hand trading platforms on continuance intention to …[J].Review,2022,2(1):1-9. doi:10.1000/example.",
        provider="serpapi",
    )

    assert records == []


def test_merge_reference_records_deduplicates_and_renumbers() -> None:
    first = [ReferenceRecord(index=1, title="文献甲", formatted="[1]作者.文献甲[J].期刊,2024.")]
    second = [
        ReferenceRecord(index=1, title="文献甲", formatted="[1]作者.文献甲[J].期刊,2024."),
        ReferenceRecord(
            index=2,
            title="文献乙",
            doi="10.1000/example",
            formatted="[2]作者.文献乙[J].期刊,2023. doi:10.1000/example.",
        ),
        ReferenceRecord(
            index=3,
            title="文献乙的近似题名",
            doi="https://doi.org/10.1000/EXAMPLE",
            formatted="[3]作者.文献乙的近似题名[J].期刊,2023. doi:10.1000/example.",
        ),
    ]
    merged = _merge_records(first, second, 10)
    assert [item.index for item in merged] == [1, 2]
    assert merged[1].formatted.startswith("[2]")


@pytest.mark.parametrize("chinese,english", [(10, 5), (0, 5), (3, 1), (10, 0)])
def test_reference_language_quota_and_selection(chinese: int, english: int) -> None:
    records = [
        ReferenceRecord(index=index, title=f"中文{index}", language="zh", formatted=f"[{index}]中文{index}")
        for index in range(1, 11)
    ] + [
        ReferenceRecord(
            index=index + 10, title=f"English {index}", language="en", formatted=f"[{index + 10}]English {index}"
        )
        for index in range(1, 6)
    ]
    selected = _select_records(records, chinese + english, chinese, english)
    assert len(selected) == chinese + english
    assert sum(item.language == "zh" for item in selected) == chinese
    assert sum(item.language == "en" for item in selected) == english
    assert [item.index for item in selected] == list(range(1, chinese + english + 1))


def test_reference_lexical_ranking_only_admits_direct_chinese_topic_match() -> None:
    def record(index: int, title: str, language: Literal["zh", "en"]) -> ReferenceRecord:
        return ReferenceRecord(
            index=index,
            title=title,
            authors=["Author A"],
            year="2024",
            source="Journal of Information Systems",
            volume="12",
            issue="2",
            pages="10-20",
            doi="10.1000/test" if language == "en" else "",
            language=language,
            formatted=f"[{index}]Author A.{title}[J].Journal of Information Systems,2024,12(2):10-20.",
        )

    records = [
        record(1, "Campus second-hand trading platform design with Spring Boot", "en"),
        record(2, "校园二手交易平台设计研究", "zh"),
        record(3, "Design and implementation of robot assisted chemistry online Q&A using Spring Boot", "en"),
        record(4, "A higher-performance big data-based movie recommendation system", "en"),
        record(5, "On-demand fashion: wardrobe management and trading community", "en"),
        record(7, "校园综合服务平台设计与实现", "zh"),
    ]
    records.append(
        record(6, "Campus second-hand marketplace research", "en").model_copy(update={"pages": "", "doi": ""})
    )
    ranked = _rank_records_by_relevance(
        "基于 Spring Boot 与 Vue 的校园二手交易平台设计与实现",
        "Spring Boot；Vue；MySQL",
        records,
    )
    titles = {item.title for item in ranked}
    assert titles == {"校园二手交易平台设计研究"}


def test_reference_relevance_does_not_treat_outline_concepts_as_topic_gates() -> None:
    def record(index: int, title: str) -> ReferenceRecord:
        return ReferenceRecord(
            index=index,
            title=title,
            authors=["测试作者"],
            year="2024",
            source="管理科学",
            volume="12",
            issue="2",
            pages="10-20",
            language="zh",
            formatted=f"[{index}]测试作者.{title}[J].管理科学,2024,12(2):10-20.",
        )

    ranked = _rank_records_by_relevance(
        "制造企业供应链一体化营销的协同机制构建",
        "研究市场竞争力、关键要素识别与信息系统建设",
        [
            record(1, "供应链协同机制构建研究"),
            record(2, "制造企业供应链整合路径研究"),
            record(3, "图像识别技术的市场应用研究"),
        ],
    )

    assert {item.title for item in ranked} == {
        "供应链协同机制构建研究",
        "制造企业供应链整合路径研究",
    }


def test_reference_validation_requires_minimum_coverage() -> None:
    result = {
        "abstract": "摘" * 220,
        "introduction": "已有研究[1]。",
        "references": [{"index": 1}, {"index": 2}],
    }
    try:
        _validate_result("literature_review", result, 2)
    except RuntimeError as exc:
        assert "正文引用参考文献不足" in str(exc)
    else:
        raise AssertionError("正文引用量不足时应校验失败")


def test_source_outline_citations_do_not_count_as_generated_body_citations() -> None:
    result = {
        "source_outline": [{"chapter": "研究背景[1]", "sections": []}],
        "thesis_config": {"aboutmsg": "请审查文献[2]"},
        "writing_outline": [{"title": "研究背景[3]", "sections": []}],
    }
    assert _missing_reference_indexes("proposal_report", result, 2) == [1, 2]


def test_task_book_no_citation_mode_removes_nested_body_indexes_only() -> None:
    result: dict[str, Any] = {
        "design_background": "课题背景[1]。",
        "module_tasks": [{"name": "资料整理[2]", "responsibilities": "说明依据[3]"}],
        "references": [{"formatted": "[1]真实资料"}],
    }
    material_generation._remove_inline_citations("task_book", result)
    assert result["design_background"] == "课题背景。"
    assert result["module_tasks"][0]["name"] == "资料整理"
    assert result["references"][0]["formatted"] == "[1]真实资料"


# 共享文献服务降级后，材料的后置引用校验也必须按实际零篇处理
@pytest.mark.parametrize("document_type", ["proposal_report", "literature_review"])
async def test_material_generation_without_references_is_explicit_draft(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
    document_type: str,
) -> None:
    """替换外部服务，保留真实长度、引用校验和 DOCX 构建。"""
    request = {
        "title": "测试材料",
        "target_word_count": 3500,
        "thesis_config": {"chinese_reference_count": 12, "english_reference_count": 0},
        "source_outline": [
            {"chapter": f"章节{index}", "sections": [{"name": "背景", "abstract": "背景", "subsections": []}]}
            for index in range(5)
        ],
    }
    plan = build_length_plan(document_type, request)
    content: dict[str, Any] = {field: "测" * value.target for field, value in plan.fields.items()}
    content["writing_outline"] = [{"title": f"章节{index}", "children": []} for index in range(5)]
    content["keywords"] = ["测试", "计划"]
    if plan.theme is not None:
        content["themes"] = [
            {"title": f"待核实方向{index}", "content": "测" * plan.theme.target} for index in range(plan.theme_count)
        ]
    reference_lookup = AsyncMock(return_value=[])
    monkeypatch.setattr(material_generation, "retrieve_reference_records", reference_lookup)
    monkeypatch.setattr(material_generation, "publish_progress", AsyncMock())
    monkeypatch.setattr(material_generation, "generate_proposal_content", AsyncMock(return_value=content))
    monkeypatch.setattr(material_generation, "generate_literature_review_content", AsyncMock(return_value=content))
    monkeypatch.setattr(
        material_generation, "get_settings", lambda: SimpleNamespace(THESIS_MATERIAL_OUTPUT_ROOT=str(tmp_path))
    )
    monkeypatch.setattr(llm_service, "_ask_text", AsyncMock(side_effect=AssertionError("合规字数不得调用模型")))
    result = await material_generation.generate_thesis_material_document(
        task_id="isolated-material",
        document_type=document_type,
        request_payload=request,
    )
    assert result["result_data"]["references"] == []
    assert result["result_data"]["reference_quality"]["status"] == "unavailable"
    assert result["result_data"]["quality_warnings"][0]["code"] == "references_unavailable"
    assert reference_lookup.await_args is not None
    assert reference_lookup.await_args.kwargs["english_reference_count"] == 0
    assert reference_lookup.await_args.kwargs["chinese_reference_count"] == 12
    with zipfile.ZipFile(result["docx_path"]) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "未检索到可用的真实文献" in xml
    assert "[1]" not in xml


@pytest.mark.asyncio
async def test_default_inline_citations_keeps_verified_reference_list(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    request = ProposalReportRequest.model_validate(
        {
            "title": "通用课题研究",
            "target_word_count": 2500,
            "source_outline": [
                {
                    "chapter": "绪论",
                    "sections": [{"name": "研究背景", "abstract": "背景", "subsections": []}],
                }
            ],
            "thesis_config": {"chinese_reference_count": 8, "english_reference_count": 0},
        }
    ).model_dump(mode="json")
    plan = build_length_plan("proposal_report", request)
    content: dict[str, Any] = {field: "测" * length.target for field, length in plan.fields.items()}
    content["research_purpose"] += "相关资料[1]。"
    content["writing_outline"] = [{"title": "绪论", "sections": [{"title": "研究背景", "subsections": []}]}]
    reference = ReferenceRecord(index=1, title="可核验研究", formatted="[1]测试作者.可核验研究[J].测试期刊,2025:1-3.")
    monkeypatch.setattr(material_generation, "retrieve_reference_records", AsyncMock(return_value=[reference]))
    monkeypatch.setattr(material_generation, "publish_progress", AsyncMock())
    monkeypatch.setattr(material_generation, "generate_proposal_content", AsyncMock(return_value=content))
    monkeypatch.setattr(
        material_generation, "get_settings", lambda: SimpleNamespace(THESIS_MATERIAL_OUTPUT_ROOT=str(tmp_path))
    )
    monkeypatch.setattr(llm_service, "_ask_text", AsyncMock(side_effect=AssertionError("合规字数不得调用模型")))

    output = await material_generation.generate_thesis_material_document(
        task_id="no-inline-cites",
        document_type="proposal_report",
        request_payload=request,
    )

    result = output["result_data"]
    assert "[1]" in result["research_purpose"]
    assert len(result["references"]) == 1
    with zipfile.ZipFile(output["docx_path"]) as archive:
        xml = archive.read("word/document.xml").decode()
    assert "可核验研究" in xml


def test_length_repair_only_regenerates_invalid_fields(monkeypatch: pytest.MonkeyPatch) -> None:
    calls: list[str] = []
    request = {"title": "测试课题", "target_word_count": 2500}
    plan = build_length_plan("proposal_report", request)

    async def fake_ask_text(system: str, prompt: str, *, max_tokens: int = 5000) -> str:
        del system, max_tokens
        calls.append(prompt)
        return "修" * plan.fields["research_purpose"].target

    monkeypatch.setattr("services.thesis_material.llm_service._ask_text", fake_ask_text)
    result: dict[str, Any] = {field: "正" * length.target for field, length in plan.fields.items()}
    result["research_purpose"] = "长" * 1200

    asyncio.run(repair_length_constraints("proposal_report", request, result))

    assert len(calls) == 1
    assert len(result["research_purpose"]) == plan.fields["research_purpose"].target
    assert len(result["research_status_and_trends"]) == plan.fields["research_status_and_trends"].target


def test_literature_minor_section_shortfall_keeps_valid_total(monkeypatch: pytest.MonkeyPatch) -> None:
    request = {"title": "测试课题", "target_word_count": 6000}
    plan = build_length_plan("literature_review", request)
    assert plan.theme is not None
    short_conclusion = "正" * round(plan.fields["conclusion"].minimum * 0.88)
    result: dict[str, Any] = {field: "正" * length.target for field, length in plan.fields.items()}
    result["conclusion"] = short_conclusion
    result["themes"] = [
        {"title": f"主题{index}", "content": "正" * plan.theme.target} for index in range(plan.theme_count)
    ]
    monkeypatch.setattr(llm_service, "_ask_text", AsyncMock(return_value=short_conclusion))

    asyncio.run(repair_length_constraints("literature_review", request, result))

    assert result["conclusion"] == short_conclusion
    _validate_result("literature_review", result, 0, request)


def test_literature_final_shortfall_uses_confirmed_outline() -> None:
    request = {
        "title": "测试课题",
        "target_word_count": 6000,
        "source_outline": [
            {"chapter": "系统需求分析"},
            {"chapter": "系统设计"},
            {"chapter": "系统测试"},
        ],
    }
    plan = build_length_plan("literature_review", request)
    assert plan.theme is not None
    result: dict[str, Any] = {field: "正" * length.minimum for field, length in plan.fields.items()}
    result["domestic_research"] = result["domestic_research"][:-220]
    result["themes"] = [
        {"title": f"主题{index}", "content": "正" * plan.theme.minimum} for index in range(plan.theme_count)
    ]

    _ensure_body_minimum_after_normalization("literature_review", request, result)

    assert _literature_body_length(result) >= plan.body_minimum
    assert "系统需求分析" in result["conclusion"]


def test_target_word_count_drives_dynamic_section_budgets(monkeypatch: pytest.MonkeyPatch) -> None:
    async def unexpected_llm_call(*_args: object, **_kwargs: object) -> str:
        raise AssertionError("处于动态预算内时不应调用模型")

    monkeypatch.setattr("services.thesis_material.llm_service._ask_text", unexpected_llm_call)
    request = {"title": "测试课题", "target_word_count": 3500}
    plan = build_length_plan("literature_review", request)
    assert plan.theme is not None
    result: dict[str, Any] = {field: "正" * length.target for field, length in plan.fields.items()}
    result["themes"] = [
        {"title": f"主题{index}", "content": "正" * plan.theme.target} for index in range(plan.theme_count)
    ]

    asyncio.run(repair_length_constraints("literature_review", request, result))

    assert plan.theme_count == 3
    assert plan.body_minimum == 3132
    assert plan.body_maximum == 3868
    assert plan.body_minimum <= _literature_body_length(result) <= plan.body_maximum
    assert 3400 <= _literature_body_length(result) <= 3600


def test_reference_coverage_repair_preserves_dynamic_theme_count(monkeypatch: pytest.MonkeyPatch) -> None:
    async def fake_ask_text(*_args: object, **_kwargs: object) -> str:
        return "补充比较[12]。"

    monkeypatch.setattr("services.thesis_material.llm_service._ask_text", fake_ask_text)
    request = {"title": "测试课题", "target_word_count": 3500}
    result = {"themes": [{"title": f"主题{index}", "content": "原有内容。"} for index in range(3)]}
    missing = [
        ReferenceRecord(
            index=12,
            language="zh",
            title="高校二手交易平台研究",
            authors=["张三"],
            year="2025",
            source="计算机应用研究",
            pages="1-8",
            formatted="[12]张三.高校二手交易平台研究[J].计算机应用研究,2025:1-8.",
        )
    ]

    asyncio.run(repair_reference_coverage("literature_review", request, result, missing))

    assert len(result["themes"]) == 3
    assert "补充比较[12]" in result["themes"][-1]["content"]


def test_section_validation_allows_small_post_processing_drift() -> None:
    request = {"target_word_count": 3500}
    plan = build_length_plan("literature_review", request)
    assert plan.theme is not None
    result: dict[str, Any] = {field: "正" * length.target for field, length in plan.fields.items()}
    result["introduction"] += "[1]"
    result["foreign_research"] = "正" * (plan.fields["foreign_research"].maximum + 20) + "[2]"
    result["themes"] = [
        {"title": f"主题{index}", "content": "正" * plan.theme.target} for index in range(plan.theme_count)
    ]

    _validate_result("literature_review", result, 2, request)


def test_word_count_metadata_excludes_outline_schedule_and_references() -> None:
    request = {"target_word_count": 2500}
    plan = build_length_plan("proposal_report", request)
    result: dict[str, Any] = {field: "正" * length.target for field, length in plan.fields.items()}
    result.update({"writing_outline": ["不计入" * 1000], "schedule": ["不计入" * 1000]})
    metadata = _word_count_metadata("proposal_report", request, result)
    assert metadata["actual"] == _material_body_char_count("proposal_report", result)
    assert metadata["target"] == 2500
    assert "writing_outline" in metadata["excluded_sections"]


def test_post_normalization_shortage_uses_uncited_topic_supplement() -> None:
    request = {"target_word_count": 3500}
    plan = build_length_plan("literature_review", request)
    assert plan.theme is not None
    result: dict[str, Any] = {field: "正" * length.target for field, length in plan.fields.items()}
    result["themes"] = [
        {"title": f"主题{index}", "content": "正" * plan.theme.target} for index in range(plan.theme_count)
    ]
    result["domestic_research"] = result["domestic_research"][:-400]

    _ensure_body_minimum_after_normalization("literature_review", request, result)

    assert _material_body_char_count("literature_review", result) >= plan.body_minimum
    assert "[" not in result["conclusion"]


def test_literature_final_shortfall_can_use_all_confirmed_sections() -> None:
    request = {
        "title": "校园二手交易平台研究",
        "target_word_count": 6000,
        "source_outline": [
            {
                "chapter": f"研究主题{index}",
                "sections": [{"name": f"研究子题{index}-{section}"} for section in range(2)],
            }
            for index in range(6)
        ],
    }
    plan = build_length_plan("literature_review", request)
    assert plan.theme is not None
    result: dict[str, Any] = {field: "正" * length.target for field, length in plan.fields.items()}
    result["themes"] = [
        {"title": f"主题{index}", "content": "正" * plan.theme.target} for index in range(plan.theme_count)
    ]
    excess = _material_body_char_count("literature_review", result) - plan.body_minimum + 450
    for field in ("domestic_research", "foreign_research", "method_comparison", "introduction"):
        removed = min(len(result[field]), excess)
        result[field] = result[field][:-removed]
        excess -= removed
        if excess <= 0:
            break
    assert _material_body_char_count("literature_review", result) == plan.body_minimum - 450

    _ensure_body_minimum_after_normalization("literature_review", request, result)

    assert _material_body_char_count("literature_review", result) >= plan.body_minimum
    assert "研究子题" in result["conclusion"]


def test_task_book_generated_numeric_metrics_are_marked_as_unconfirmed() -> None:
    result = {
        "main_indicators": [
            "接口平均响应时间不超过500ms",
            "支持100名并发用户",
            "完成核心业务流程核验",
            "异常操作应有明确提示",
        ]
    }
    _mark_unconfirmed_generated_metrics({"research_context": {}}, result)
    assert result["generated_suggestion_fields"] == ["main_indicators"]
    assert result["main_indicators"][0].startswith("建议值（待导师确认）：")
    assert result["main_indicators"][2] == "完成核心业务流程核验"

    already_marked = {"main_indicators": ["单张图片建议值5MB，待导师确认"] * 4}
    _mark_unconfirmed_generated_metrics({"research_context": {}}, already_marked)
    assert already_marked["generated_suggestion_fields"] == ["main_indicators"]
    assert already_marked["main_indicators"][0] == "单张图片建议值5MB，待导师确认"

    all_fields = {
        "design_goals": ["完成核心流程"] * 5,
        "deliverable_requirements": ["演示视频时长5-10分钟"] * 2,
        "main_indicators": ["功能可验证"] * 4,
    }
    _mark_unconfirmed_generated_metrics({"research_context": {}}, all_fields)
    assert all_fields["generated_suggestion_fields"] == ["deliverable_requirements"]
    assert all_fields["deliverable_requirements"][0].startswith("建议值（待导师确认）：")

    partly_confirmed = {"main_indicators": ["支持并发50人", "接口响应不超过300ms"]}
    _mark_unconfirmed_generated_metrics(
        {"thesis_config": {"aboutmsg": "导师要求并发50人"}},
        partly_confirmed,
    )
    assert partly_confirmed["main_indicators"][0] == "支持并发50人"
    assert partly_confirmed["main_indicators"][1].startswith("建议值（待导师确认）：")


def test_unsupported_citation_conclusion_is_rewritten_from_reference_title() -> None:
    result = {
        "foreign_research": "某学者[1]发现环保意识显著提高购买意愿。另一研究[2]围绕平台架构展开。",
        "references": [
            {"index": 1, "title": "Second-hand clothing shopping among college students"},
            {"index": 2, "title": "Campus marketplace architecture"},
        ],
    }
    normalize_citation_claims(result)
    assert "显著提高" not in result["foreign_research"]
    assert "Second-hand clothing shopping among college students" in result["foreign_research"]
    assert "另一研究[2]围绕平台架构展开" in result["foreign_research"]


def test_repeated_unsupported_citation_claims_are_collapsed() -> None:
    result = {
        "foreign_research": (
            "研究[1]表明效果显著。研究[1]证明效率提高。研究[2]显示体验改善。文献[2]采用问卷分析用户行为。"
        ),
        "references": [
            {"index": 1, "title": "Campus marketplace architecture"},
            {"index": 2, "title": "College second-hand trading"},
        ],
    }
    normalize_citation_claims(result)
    assert result["foreign_research"].count("文献[1]") == 1
    assert result["foreign_research"].count("文献[2]") == 1
    assert "采用问卷" not in result["foreign_research"]


def test_sentence_trim_preserves_citations_and_required_topics() -> None:
    text = "".join(
        (
            "行业背景" + "甲" * 120 + "。",
            "现实问题" + "乙" * 120 + "。",
            "技术背景" + "丙" * 120 + "。",
            "研究必要性" + "丁" * 120 + "。",
            "应用价值" + "戊" * 120 + "。",
            "研究目标" + "己" * 120 + "。",
            "补充分析" + "庚" * 120 + "。",
            "相关依据[1]" + "辛" * 120 + "。",
            "重复说明" + "壬" * 120 + "。",
        )
    )

    trimmed = _trim_to_complete_sentences(text, "research_purpose", 700, 1000)

    assert 700 <= len(trimmed) <= 1000
    assert "[1]" in trimmed
    for keyword in ("背景", "问题", "技术", "必要性", "价值", "目标"):
        assert keyword in trimmed


def test_character_limit_is_deterministic_for_uncited_short_section() -> None:
    trimmed = _trim_to_character_limit("关键业务与数据设计，" * 80, 239)
    assert len(trimmed) == 239
    assert trimmed.endswith("。")


def test_relative_schedule_has_required_number_of_stages() -> None:
    proposal = _build_schedule(PROPOSAL_SCHEDULE, default_weeks=16)
    task_book = _build_schedule(TASK_BOOK_SCHEDULE, default_weeks=20)
    assert len(proposal) == 9
    assert proposal[0]["start"] == "第1周"
    assert proposal[-1]["end"] == "第16周"
    assert len(task_book) == 5


def test_build_three_document_types(tmp_path: Path) -> None:
    references = [
        ReferenceRecord(
            index=1,
            title="测试文献",
            authors=["张三"],
            year="2024",
            formatted="[1]张三.测试文献[J].测试期刊,2024,1(1):1-5.",
        )
    ]
    request: dict[str, Any] = {"research_context": {}}
    proposal_result = {
        "research_purpose": "研究目的[1]。",
        "research_status_and_trends": "研究现状[1]。",
        "key_points": "研究重点。",
        "difficulties": "研究难点。",
        "research_methods": "研究方法。",
        "schedule": _build_schedule(PROPOSAL_SCHEDULE, default_weeks=16),
    }
    review_result = {
        "abstract": "摘要。",
        "keywords": ["测试", "综述", "方法"],
        "introduction": "引言[1]。",
        "domestic_research": "国内研究[1]。",
        "foreign_research": "国外研究[1]。",
        "themes": [{"title": "主题", "content": "比较研究[1]。"}] * 3,
        "method_comparison": "方法比较。",
        "research_gaps": "不足。",
        "future_trends": "趋势。",
        "conclusion": "结论。",
    }
    task_result = {
        "design_background": "设计背景。",
        "technology_stack": ["Java", "Vue"],
        "design_goals": [f"目标{i}" for i in range(1, 6)],
        "module_tasks": [
            {"name": f"模块{i}", "role": "用户", "responsibilities": "完成业务", "boundary": "仅处理本模块"}
            for i in range(1, 5)
        ],
        "schedule_items": _build_schedule(TASK_BOOK_SCHEDULE, default_weeks=20),
        "deliverable_forms": ["成果文档", "项目源文件"],
        "deliverable_requirements": ["项目可运行", "测试通过"],
    }
    paths = []
    for document_type, result in (
        ("proposal_report", proposal_result),
        ("literature_review", review_result),
        ("task_book", task_result),
    ):
        path = tmp_path / f"{document_type}.docx"
        build_thesis_material_document(
            document_type=document_type,
            title="测试课题",
            request=request,
            result=result,
            references=references if document_type != "task_book" else [],
            output_path=path,
        )
        paths.append(path)
    assert all(path.exists() for path in paths)
    for path in (paths[0], paths[2]):
        document = Document(str(path))
        visible_text = "\n".join(
            [paragraph.text for paragraph in document.paragraphs]
            + [cell.text for table in document.tables for row in table.rows for cell in row.cells]
        )
        assert "【待补充：学校】" in visible_text
        assert "【待补充：学生姓名】" in visible_text
    for path in paths:
        with zipfile.ZipFile(path) as package:
            package_xml = "\n".join(
                package.read(name).decode("utf-8") for name in ("word/document.xml", "word/styles.xml")
            )
        assert "STHeiti" not in package_xml
        assert 'w:ascii="Times New Roman"' in package_xml
        assert 'w:hAnsi="Times New Roman"' in package_xml
        assert 'w:eastAsia="宋体"' in package_xml
        assert 'w:eastAsia="黑体"' in package_xml
    assert "指导教师意见" in "\n".join(
        cell.text for table in Document(str(paths[0])).tables for row in table.rows for cell in row.cells
    )
    assert "审核意见" in "\n".join(
        cell.text for table in Document(str(paths[2])).tables for row in table.rows for cell in row.cells
    )
    task_document = Document(str(paths[2]))
    task_paragraph = task_document.tables[-1].cell(0, 0).paragraphs[0]
    assert round(float(task_paragraph.paragraph_format.line_spacing), 2) == 1.15
    assert task_paragraph.runs[0].font.size.pt == 10


def test_proposal_uses_short_non_splitting_rows_and_grouped_approval(tmp_path: Path) -> None:
    result = {
        "research_purpose": "研究背景与目的。" * 120,
        "research_status_and_trends": "已有研究及其不足[1]。" * 160,
        "research_content": "主要研究内容。" * 80,
        "key_points": "研究重点。" * 40,
        "difficulties": "研究难点。" * 40,
        "research_methods": "研究方法。" * 50,
        "feasibility_and_innovation": "可行性与创新点。" * 35,
        "writing_outline": [
            {
                "title": "绪论",
                "sections": [{"title": "研究背景", "subsections": ["校园闲置物品现状"]}],
            }
        ],
        "schedule": [],
    }
    reference = ReferenceRecord(
        index=1,
        title="校园二手交易平台研究",
        authors=["张三"],
        year="2024",
        source="软件导刊",
        pages="1-8",
        formatted="[1]张三.校园二手交易平台研究[J].软件导刊,2024:1-8.",
    )
    path = tmp_path / "proposal-pagination.docx"
    build_thesis_material_document(
        document_type="proposal_report",
        title="测试课题",
        request={},
        result=result,
        references=[reference],
        output_path=path,
    )
    document = Document(str(path))
    assert len(document.tables) == 9
    assert all("w:cantSplit" in row._tr.xml for table in document.tables for row in table.rows)
    assert max(len(row.cells[0].text) for table in document.tables for row in table.rows) <= 450
    outline_paragraph = document.tables[5].cell(1, 0).paragraphs[0]
    assert outline_paragraph.paragraph_format.line_spacing == 1.0
    assert outline_paragraph.runs[0].font.size.pt == 9.5
    approval_text = document.tables[-1].cell(0, 0).text
    assert "指导教师意见" in approval_text
    assert "负责人（签章）" in approval_text
    assert "某某" not in "\n".join(paragraph.text for paragraph in document.paragraphs)
