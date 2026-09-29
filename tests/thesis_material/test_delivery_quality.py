"""真实回归发现的材料交付质量问题。"""

from pathlib import Path
from typing import Any
from unittest.mock import AsyncMock

import pytest
from docx import Document
from docx.oxml.ns import qn

from services.thesis.content.reference_service_serpapi import _resolve_doc_marker
from services.thesis_material import llm_service
from services.thesis_material.document_builder import _schedule_text, build_thesis_material_document
from services.thesis_material.llm_service import _clean_task_instructions, normalize_citation_claims


# 单个字段预算与总正文验收分离，避免局部短几十字中断整份文档。
async def test_short_field_uses_final_validation_floor(monkeypatch: pytest.MonkeyPatch) -> None:
    """117 字的有效短字段交由后续总字数校验，而非局部抛错。"""
    monkeypatch.setattr(llm_service, "_ask_text", AsyncMock(return_value="甲" * 116 + "。"))
    text = await llm_service._repair_length_value("feasibility_and_innovation", "甲" * 117, 155, 195, {"title": "测试"}, {})
    assert len(text) == 117


# 引用安全处理是字数计算的一部分。
async def test_length_repair_counts_normalized_citation_text(monkeypatch: pytest.MonkeyPatch) -> None:
    """模型短句经引用规范化扩写后必须重新计量并修复。"""
    safe_text = "相关主题仍需进一步核对原文，" * 8 + "请确认。"
    rewrite = AsyncMock(side_effect=["文献[1]证明效率提高。", safe_text])
    monkeypatch.setattr(llm_service, "_ask_text", rewrite)
    result = {"references": [{"index": 1, "title": "长题名" * 80}]}
    value = await llm_service._repair_length_value("future_trends", "短句。", 80, 150, {"title": "测试"}, result)
    assert value == safe_text
    assert rewrite.await_count == 2


# 内部占位规则不是用户任务，业务隐私保护要求仍须保留。
def test_task_internal_instructions_are_removed_without_losing_business_privacy() -> None:
    """剔除模型误抄的指令并保留合法任务内容。"""
    assert _clean_task_instructions([
        "不得填写或推断姓名、学号，相关字段由文档层固定占位。",
        "保护交易用户个人信息。",
        {"boundary": "限于校园交易；不复制论文目录作为任务书内容。"},
    ]) == ["保护交易用户个人信息。", {"boundary": "限于校园交易；"}]


# 来源类型不能被非空会议/图书容器名称覆盖。
@pytest.mark.parametrize(("source_type", "expected"), [("proceedings-article", "C"), ("book-chapter", "M"), ("journal-article", "J")])
def test_explicit_reference_type_precedes_container(source_type: str, expected: str) -> None:
    """明确的来源类型优先于容器名称。"""
    assert _resolve_doc_marker("Example proceedings", source_type, "") == expected


# 引用降级不得删除枚举项或在反复规范化中累积重复说明。
def test_citation_normalization_preserves_sequence_and_is_idempotent() -> None:
    """保留三项枚举的位置、段落和后续说明。"""
    result: dict[str, Any] = {
        "research_purpose": "一是文献[1]证明效率提高。二是文献[2]发现体验改善。三是制定研究计划。\n下一阶段核对资料。",
        "references": [{"index": 1, "title": "主题甲"}, {"index": 2, "title": "主题乙"}],
    }
    normalize_citation_claims(result)
    text = result["research_purpose"]
    assert text.index("一是") < text.index("二是") < text.index("三是")
    assert "提高" not in text and "改善" not in text
    assert "\n下一阶段" in text
    normalize_citation_claims(result)
    assert result["research_purpose"] == text


# 枚举项互相依赖，不能为满足预算删掉其中一项。
def test_sentence_trimming_keeps_dependent_enumeration() -> None:
    """不足以容纳完整枚举时交由重写，不机械删除某个编号。"""
    text = "一是" + "甲" * 60 + "。二是" + "乙" * 60 + "。三是" + "丙" * 60 + "。"
    assert llm_service._trim_to_complete_sentences(text, "research_purpose", 80, 150) == text


# 单周计划不显示同一周两次。
def test_single_week_schedule_is_compact() -> None:
    """只合并相同起止周，不改变任务内容。"""
    assert _schedule_text([{"start": "第5周", "end": "第5周", "task": "检查", "deliverable": "记录"}]).startswith("第5周：")


# 用最小任务书检查固定列宽、重复表头和真实页码域。
def test_task_book_layout_and_task_labels(tmp_path: Path) -> None:
    """生成产物应保留任务目标语义，内容列大于标签列。"""
    path = tmp_path / "task.docx"
    build_thesis_material_document(
        document_type="task_book", title="测试课题", request={}, references=[], output_path=path,
        result={"module_tasks": [{"name": "分析", "role": "确定研究范围。", "responsibilities": "梳理资料。", "boundary": "仅限公开资料。"}]},
    )
    doc = Document(str(path))
    assert doc.tables[-1].columns[1].width > doc.tables[-1].columns[0].width
    assert doc.tables[1].rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None
    assert "PAGE" in doc.sections[0].footer._element.xml
    body = "".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    assert "面向确定" not in body and "。，" not in body
