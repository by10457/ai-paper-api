import tempfile
import zipfile
from pathlib import Path

from services.thesis.document.docx_builder import build_word_document


def test_cover_main_title_does_not_use_body_exact_line_spacing() -> None:
    out = Path(tempfile.mktemp(suffix=".docx"))
    try:
        build_word_document(
            title="新媒体语境下非物质文化遗产短视频传播策略研究",
            full_text="# 第一章 绪论\n正文。\n",
            output_path=str(out),
            abstract_zh="中文摘要",
            keywords_zh="关键词",
            abstract_en="English abstract",
            keywords_en="keywords",
            acknowledgment="感谢。",
            references="[1] 文献。",
            placeholders=[],
            image_paths={},
        )

        with zipfile.ZipFile(out) as zf:
            xml = zf.read("word/document.xml").decode("utf-8", errors="ignore")

        idx = xml.find("学士学位论文")
        assert idx != -1
        window = xml[max(0, idx - 500): idx + 200]
        assert 'w:lineRule="exact"' not in window
    finally:
        out.unlink(missing_ok=True)

# 文档仅使用固定个人信息占位，元数据不得填入作者身份
def test_document_profile_placeholders_and_metadata(tmp_path: Path) -> None:
    """检查实际 DOCX 的封面、声明页与文件属性。

    Args:
        tmp_path: pytest 隔离输出目录。
    """
    from docx import Document

    path = tmp_path / "privacy.docx"
    build_word_document(
        full_text="# 1 绪论\n正文。",
        placeholders=[],
        image_paths={},
        output_path=str(path),
        title="隐私边界测试论文",
    )
    document = Document(str(path))
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
    for label in ("作者姓名", "指导教师", "专业名称", "学院（系）", "学号", "班级", "学位类别", "年月"):
        assert f"【待补充：{label}】" in xml
    assert document.core_properties.author == "AI Paper"
    assert document.core_properties.last_modified_by == "AI Paper"
