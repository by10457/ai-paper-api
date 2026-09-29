"""正文补写和交付分页的回归测试。"""

from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from docx import Document
from docx.oxml.ns import qn

from services.thesis.document.inline import _add_table
from services.thesis.document.pages import _balanced_cover_title
from services.thesis.generation import pipeline


# 安全清理后短文只能补写现有小节，不重生成原文。
async def test_short_fulltext_appends_to_existing_section(monkeypatch: pytest.MonkeyPatch) -> None:
    """保留原有内容和标题，补写后重新按可见字数检查。"""
    original = "# 1 研究方案\n## 1.1 方案说明\n" + "研究方案需要确认。" * 12
    model = SimpleNamespace(
        ainvoke=AsyncMock(return_value=SimpleNamespace(content="拟进一步核对资料来源并说明研究边界。" * 15))
    )
    monkeypatch.setattr(pipeline, "create_configured_llm", AsyncMock(return_value=model))
    result = await pipeline._repair_short_fulltext(
        original,
        title="测试课题",
        target_word_count=400,
        writing_requirements="",
        confirmed_technologies=set(),
        evidence_instruction="没有实测数据",
    )
    assert original in result
    assert result.count("## 1.1") == 1
    assert pipeline.count_visible_words(result) >= 360


# 生成器偏离格式时不能把新增章节或图表注入已有大纲。
async def test_short_fulltext_rejects_structural_additions(monkeypatch: pytest.MonkeyPatch) -> None:
    """三轮无有效补写时保留原文，最终质量校验仍可拒绝。"""
    model = SimpleNamespace(ainvoke=AsyncMock(return_value=SimpleNamespace(content="# 新增章节\n未经允许的内容")))
    monkeypatch.setattr(pipeline, "create_configured_llm", AsyncMock(return_value=model))
    original = "# 1 研究方案\n原文。"
    result = await pipeline._repair_short_fulltext(
        original,
        title="测试课题",
        target_word_count=400,
        writing_requirements="",
        confirmed_technologies=set(),
        evidence_instruction="没有实测数据",
    )
    assert result == original
    assert model.ainvoke.await_count == 3


# 跨页表头和段落分页标记要写入实际 DOCX。
def test_body_table_repeats_header(tmp_path: Path) -> None:
    """通过序列化后的 OOXML 验证表头重复，不访问外部渲染器。"""
    document = Document()
    _add_table(document, [["编号", "研究任务"], ["1", "核对资料"]])
    path = tmp_path / "table.docx"
    document.save(str(path))
    table = Document(str(path)).tables[0]
    assert table.rows[0]._tr.find(".//" + qn("w:tblHeader")) is not None
    assert table.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next


# 长封面标题应均衡断行，不拆分英文技术词。
def test_cover_title_balances_lines() -> None:
    """标题内容与单词完整保留，避免一字尾行。"""
    title = "基于Spring Boot与Vue的校园二手交易平台设计与实现"
    result = _balanced_cover_title(title, 400)
    assert result.replace("\n", "") == title
    assert min(len(line) for line in result.splitlines()) > 5
    assert "Spring" in result and "Boot" in result
