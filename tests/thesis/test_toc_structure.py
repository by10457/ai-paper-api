"""Tests for the visible TOC + PAGEREF implementation in docx_builder.

Each test generates a minimal DOCX via build_word_document(), unzips it,
and inspects the raw word/document.xml to verify the OOXML structure.
"""

import re
import tempfile
import zipfile
from collections.abc import Iterator
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from docx import Document as DocumentFactory

from services.thesis.document.docx_builder import (
    _pre_scan_headings,
    build_word_document,
)

# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

# Minimal body text covering all 3 heading levels
SAMPLE_BODY = """\
# 1 绪论
## 1.1 研究背景
正文段落。
## 1.2 研究意义
正文段落。
# 2 相关技术
## 2.1 Spring Boot 框架概述
### 2.1.1 核心特性
正文段落。
## 2.2 MySQL 数据库
正文段落。
# 3 系统设计
## 3.1 总体架构设计
正文段落。
"""

SAMPLE_TITLE = "基于 Spring Boot 的校园管理系统"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def _build_sample_docx() -> Path:
    """Generate a minimal DOCX and return its path."""
    out = Path(tempfile.mktemp(suffix=".docx"))
    build_word_document(
        title=SAMPLE_TITLE,
        full_text=SAMPLE_BODY,
        output_path=str(out),
        author="测试用户",
        advisor="导师",
        degree_type="学士",
        major="软件工程",
        school="计算机学院",
        year_month="2026年5月",
        abstract_zh="中文摘要内容。",
        keywords_zh="关键词1；关键词2",
        abstract_en="English abstract content.",
        keywords_en="keyword1; keyword2",
        acknowledgment="感谢所有帮助过我的人。",
        references="[1] 某参考文献。",
        placeholders=[],
        image_paths={},
    )
    return out


def _read_document_xml(docx_path: Path) -> str:
    """Extract word/document.xml as a UTF-8 string."""
    with zipfile.ZipFile(docx_path) as zf:
        return zf.read("word/document.xml").decode("utf-8")


# 读取目录 PAGEREF 域当前保存的缓存页码
def _read_cached_pageref_pages(document_xml: str) -> dict[str, int]:
    """读取目录字段在 Word 首次排版前显示的缓存页码。

    Args:
        document_xml: DOCX 中的 ``word/document.xml`` 内容。

    Returns:
        以书签名为键、缓存页码为值的映射。
    """
    pages: dict[str, int] = {}
    tree = ET.fromstring(document_xml)
    for paragraph in tree.findall(".//w:p", NS):
        instruction = "".join(node.text or "" for node in paragraph.findall(".//w:instrText", NS))
        match = re.search(r"PAGEREF\s+(_toc_\d+)", instruction)
        if match is None:
            continue
        result_nodes = paragraph.findall(".//w:t", NS)
        if result_nodes and (result_nodes[-1].text or "").isdigit():
            pages[match.group(1)] = int(result_nodes[-1].text or "0")
    return pages


@pytest.fixture(scope="module")
def generated_docx() -> Iterator[Path]:
    """Module-scoped fixture: generate once, reuse across tests."""
    path = _build_sample_docx()
    yield path
    path.unlink(missing_ok=True)


@pytest.fixture(scope="module")
def document_xml(generated_docx: Path) -> str:
    return _read_document_xml(generated_docx)


# ---------------------------------------------------------------------------
# 1. No legacy TOC field code
# ---------------------------------------------------------------------------


class TestNoLegacyTOC:
    """Ensure the old TOC \\o field and hint text are gone."""

    def test_no_toc_field_instruction(self, document_xml: str) -> None:
        """document.xml must NOT contain 'TOC \\o' field instruction."""
        assert "TOC \\o" not in document_xml
        assert "TOC \\\\o" not in document_xml

    def test_no_hint_text(self, document_xml: str) -> None:
        """The old '目录生成完毕' prompt text must be absent."""
        assert "目录生成完毕" not in document_xml
        assert "按 F9 更新" not in document_xml


# ---------------------------------------------------------------------------
# 2. Bookmarks on body headings
# ---------------------------------------------------------------------------


class TestBookmarks:
    """Bookmarks _toc_0, _toc_1, ... must exist on body headings."""

    def test_bookmarks_present(self, document_xml: str) -> None:
        tree = ET.fromstring(document_xml)
        bookmarks = tree.findall(".//w:bookmarkStart", NS)
        toc_bookmarks = [
            bm.get(f"{{{NS['w']}}}name") for bm in bookmarks if (bm.get(f"{{{NS['w']}}}name") or "").startswith("_toc_")
        ]
        assert len(toc_bookmarks) > 0, "No _toc_* bookmarks found"

    def test_bookmark_count_matches_headings(self, document_xml: str) -> None:
        """Number of _toc_* bookmarks should equal number of TOC entries."""
        expected = _pre_scan_headings(SAMPLE_BODY, title=SAMPLE_TITLE)

        tree = ET.fromstring(document_xml)
        bookmarks = tree.findall(".//w:bookmarkStart", NS)
        toc_bookmarks = [bm for bm in bookmarks if (bm.get(f"{{{NS['w']}}}name") or "").startswith("_toc_")]
        assert len(toc_bookmarks) == len(expected), f"Expected {len(expected)} bookmarks, got {len(toc_bookmarks)}"


# ---------------------------------------------------------------------------
# 3. PAGEREF fields in TOC page
# ---------------------------------------------------------------------------


class TestPagerefFields:
    """PAGEREF _toc_N \\h fields must exist in the document."""

    def test_pageref_present(self, document_xml: str) -> None:
        matches = re.findall(r"PAGEREF\s+_toc_\d+", document_xml)
        assert len(matches) > 0, "No PAGEREF _toc_* fields found"

    def test_pageref_count_matches_entries(self, document_xml: str) -> None:
        expected = _pre_scan_headings(SAMPLE_BODY, title=SAMPLE_TITLE)
        matches = re.findall(r"PAGEREF\s+_toc_\d+", document_xml)
        assert len(matches) == len(expected), f"Expected {len(expected)} PAGEREF fields, got {len(matches)}"

    def test_back_matter_cached_pages_follow_body(self, document_xml: str) -> None:
        """参考文献和致谢的缓存页码应位于正文之后并保持先后顺序。"""
        entries = _pre_scan_headings(SAMPLE_BODY, title=SAMPLE_TITLE)
        bookmarks = {str(entry["text"]): str(entry["bookmark"]) for entry in entries}
        cached_pages = _read_cached_pageref_pages(document_xml)

        reference_page = cached_pages[bookmarks["参考文献"]]
        acknowledgment_page = cached_pages[bookmarks["致谢"]]

        assert reference_page > 1
        assert acknowledgment_page > reference_page

    def test_long_references_push_acknowledgment_to_later_page(self, tmp_path: Path) -> None:
        """多页参考文献应继续推后致谢的目录缓存页码。"""
        references = "\n".join(
            f"[{index}] 测试作者.乡村数字治理与公共服务协同机制研究的虚构参考文献条目"
            "及其应用成效分析[J].测试学报,2026,12(3):100-120."
            for index in range(1, 21)
        )
        output_path = tmp_path / "long-references.docx"
        build_word_document(
            title=SAMPLE_TITLE,
            full_text=SAMPLE_BODY,
            output_path=str(output_path),
            acknowledgment="感谢所有提供帮助的老师和同学。",
            references=references,
            placeholders=[],
            image_paths={},
        )
        document_xml = _read_document_xml(output_path)
        entries = _pre_scan_headings(SAMPLE_BODY, title=SAMPLE_TITLE)
        bookmarks = {str(entry["text"]): str(entry["bookmark"]) for entry in entries}
        cached_pages = _read_cached_pageref_pages(document_xml)

        reference_page = cached_pages[bookmarks["参考文献"]]
        acknowledgment_page = cached_pages[bookmarks["致谢"]]

        assert acknowledgment_page >= reference_page + 2


# ---------------------------------------------------------------------------
# 4. TOC entry count matches body heading count
# ---------------------------------------------------------------------------


class TestTocEntryCount:
    """The visible TOC should have exactly as many entries as body headings."""

    def test_entry_count(self, document_xml: str) -> None:
        expected = _pre_scan_headings(SAMPLE_BODY, title=SAMPLE_TITLE)
        # Each TOC entry has a tab character followed by PAGEREF
        # Count PAGEREF occurrences as proxy for visible entries
        matches = re.findall(r"PAGEREF\s+_toc_\d+", document_xml)
        assert len(matches) == len(expected)


# ---------------------------------------------------------------------------
# 5. Non-body heading blacklist
# ---------------------------------------------------------------------------


class TestBlacklist:
    """_pre_scan_headings must filter out non-body headings."""

    @pytest.mark.parametrize(
        "heading",
        [
            "摘要",
            "摘 要",
            "中文摘要",
            "Abstract",
            "abstract",
            "ABSTRACT",
            "致谢",
            "致 谢",
            "参考文献",
        ],
    )
    def test_blacklisted_headings_excluded(self, heading: str) -> None:
        body = f"# {heading}\n正文内容\n# 第一章 绪论\n正文\n"
        entries = _pre_scan_headings(body, title="某论文题目", include_back_matter=False)
        texts = [e["text"] for e in entries]
        assert heading not in texts, f"'{heading}' should be filtered out"

    def test_title_excluded(self) -> None:
        body = "# 我的毕业论文\n# 第一章 绪论\n正文\n"
        entries = _pre_scan_headings(body, title="我的毕业论文")
        texts = [e["text"] for e in entries]
        assert "我的毕业论文" not in texts

    def test_normal_headings_kept(self) -> None:
        body = "# 1 绪论\n## 1.1 研究背景\n正文\n"
        entries = _pre_scan_headings(body, title="某论文")
        assert len(entries) == 4
        assert entries[0]["text"] == "1 绪论"
        assert entries[1]["text"] == "1.1 研究背景"
        assert entries[-2]["text"] == "参考文献"
        assert entries[-1]["text"] == "致谢"

    def test_headings_without_markdown_space_are_kept(self) -> None:
        """模型省略井号后的空格时仍应识别目录标题。"""

        body = "#1 绪论\n##1.1 研究背景\n###1.1.1 研究对象\n正文\n"

        entries = _pre_scan_headings(body, title="某论文", include_back_matter=False)

        assert [(entry["level"], entry["text"]) for entry in entries] == [
            (1, "1 绪论"),
            (2, "1.1 研究背景"),
            (3, "1.1.1 研究对象"),
        ]


def test_headings_inside_code_fences_are_excluded(tmp_path: Path) -> None:
    """代码注释不能污染预扫描目录或 Word 正文标题。"""

    body = """\
# 1 绪论
正文内容。
```python
# 数据加载
## 缺失值处理
print("hello")
```
# 2 系统实现
正文内容。
"""

    entries = _pre_scan_headings(body, include_back_matter=False)
    assert [(entry["level"], entry["text"]) for entry in entries] == [
        (1, "1 绪论"),
        (1, "2 系统实现"),
    ]

    output_path = tmp_path / "code-fence.docx"
    build_word_document(
        title="代码块目录测试",
        full_text=body,
        output_path=str(output_path),
        placeholders=[],
        image_paths={},
    )
    document = DocumentFactory(str(output_path))
    heading_texts = [
        paragraph.text
        for paragraph in document.paragraphs
        if (paragraph.style.name if paragraph.style is not None else "").startswith("Heading")
    ]
    assert "1 绪论" in heading_texts
    assert "2 系统实现" in heading_texts
    assert "数据加载" not in heading_texts
    assert "缺失值处理" not in heading_texts
    assert any(paragraph.text == "# 数据加载" for paragraph in document.paragraphs)
